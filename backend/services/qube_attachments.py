"""QUBE 会话附件 — 上传解析 / 文本提取 / 规范表格 CSV / LLM 清单构建

文件布局（data/ 为 gitignore 的本地数据目录）：

    data/qube_attachments/<session_id>/<attachment_id>_<filename>     原始文件
    data/qube_attachments/<session_id>/<attachment_id>.extract.txt    提取全文（read_attachment 分页读取）
    data/qube_attachments/<session_id>/<attachment_id>.table.csv      表格类规范 CSV（沙箱研究代码）

kind 两类：text（pdf/docx/pptx/txt/md）与 table（csv/xlsx）。表格类额外产出规范
CSV 并记录 shape/列名/head 预览，供 run_research_code 注入 attachments 变量。
"""

from __future__ import annotations

import io
import json
import re
import uuid
from pathlib import Path

import pandas as pd
from loguru import logger

from backend.config import settings

MAX_FILE_SIZE = 25 * 1024 * 1024  # 单文件 25MiB（对齐参考站文档上限）
MAX_TEXT_CHARS = 400_000  # 提取全文上限，超出截断并标注
PREVIEW_CHARS = 1_200  # 清单里的开头预览长度
TABLE_PREVIEW_ROWS = 5

TEXT_SUFFIXES = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown"}
TABLE_SUFFIXES = {".csv", ".xlsx"}

_TABLE_TRUNCATED_NOTE = "[表格过大，预览截断]"


def attachment_dir(session_id: str) -> Path:
    return Path(settings.data_dir) / "qube_attachments" / session_id


def _safe_filename(name: str) -> str:
    name = Path(name or "attachment").name
    name = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "_", name).strip(". ")
    return name or "attachment"


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"## 第 {i} 页\n{text}")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
        if texts:
            slides.append(f"## 第 {i} 页\n" + "\n\n".join(texts))
    return "\n\n".join(slides)


def _raw_csv_preview(data: bytes) -> str:
    """CSV 预览直接取原始文本行（避免 pandas 类型推断改变展示，如代码列前导零）"""
    lines = data.decode("utf-8", errors="replace").splitlines()[: TABLE_PREVIEW_ROWS + 1]
    text = "\n".join(lines)
    if len(text) > PREVIEW_CHARS:
        return text[:PREVIEW_CHARS] + _TABLE_TRUNCATED_NOTE
    return text


def _table_preview(df: pd.DataFrame) -> str:
    cols = ", ".join(str(c) for c in df.columns[:20])
    head = df.head(TABLE_PREVIEW_ROWS).to_csv(index=False)
    preview = f"{len(df)} 行 × {len(df.columns)} 列；列: {cols}\n前 {TABLE_PREVIEW_ROWS} 行:\n{head}"
    if len(preview) > PREVIEW_CHARS:
        return preview[:PREVIEW_CHARS] + _TABLE_TRUNCATED_NOTE
    return preview


def parse_upload(session_id: str, filename: str, data: bytes) -> dict:
    """解析并落盘一个上传附件，返回 qube_attachments 行 dict；失败抛 ValueError（面向用户）"""
    filename = _safe_filename(filename)
    if len(data) > MAX_FILE_SIZE:
        raise ValueError(f"文件过大（{len(data) / 1024 / 1024:.1f}MB > 25MB 上限）")
    if not data:
        raise ValueError("文件内容为空")

    suffix = Path(filename).suffix.lower()
    att_id = f"att_{uuid.uuid4().hex[:12]}"
    directory = attachment_dir(session_id)
    directory.mkdir(parents=True, exist_ok=True)
    raw_path = directory / f"{att_id}_{filename}"
    raw_path.write_bytes(data)

    kind = "table" if suffix in TABLE_SUFFIXES else "text"
    if suffix not in TEXT_SUFFIXES and suffix not in TABLE_SUFFIXES:
        # 未知扩展名按纯文本尝试（宽容处理 .log/.dat 等文本数据）
        kind = "text"

    table_meta: dict = {}
    extracted_chars = 0
    preview = ""
    try:
        if kind == "table":
            if suffix == ".csv":
                df = pd.read_csv(io.BytesIO(data))
                # 原样保留原始字节作为规范 CSV：pandas 重序列化会吃掉代码列前导零
                csv_path = directory / f"{att_id}.table.csv"
                csv_path.write_bytes(data)
                preview = _raw_csv_preview(data)
            elif suffix == ".xlsx":
                df = pd.read_excel(io.BytesIO(data), engine="openpyxl")
                csv_path = directory / f"{att_id}.table.csv"
                df.to_csv(csv_path, index=False)
                preview = _table_preview(df)
            else:  # pragma: no cover — 上方已归入 text
                raise ValueError(f"不支持的表格格式: {suffix}")
            table_meta = {
                "rows": len(df),
                "cols": len(df.columns),
                "columns": [str(c) for c in df.columns],
            }
            extracted_chars = len(csv_path.read_text(encoding="utf-8"))
        else:
            if suffix == ".pdf":
                text = _extract_pdf(data)
            elif suffix == ".docx":
                text = _extract_docx(data)
            elif suffix == ".pptx":
                text = _extract_pptx(data)
            else:
                text = data.decode("utf-8", errors="replace")
            original_chars = len(text)
            if original_chars > MAX_TEXT_CHARS:
                text = text[:MAX_TEXT_CHARS] + (
                    f"\n\n[文档过长，已截断，原始共 {original_chars} 字符]"
                )
            extracted_chars = len(text)
            preview = text[:PREVIEW_CHARS]
            (directory / f"{att_id}.extract.txt").write_text(text, encoding="utf-8")
    except ValueError:
        raise
    except Exception as e:
        logger.warning(f"附件解析失败 {filename}: {e}")
        raise ValueError(f"解析失败（{type(e).__name__}）: {e}") from e

    return {
        "id": att_id,
        "session_id": session_id,
        "filename": filename,
        "kind": kind,
        "size": len(data),
        "extracted_chars": extracted_chars,
        "preview": preview,
        "table_meta_json": json.dumps(table_meta, ensure_ascii=False),
        "raw_path": str(raw_path),
    }


def load_extract_text(att: dict) -> str:
    """读取附件提取全文（表格类读规范 CSV 文本）"""
    directory = attachment_dir(att["session_id"])
    suffix = ".table.csv" if att["kind"] == "table" else ".extract.txt"
    path = directory / f"{att['id']}{suffix}"
    try:
        return path.read_text(encoding="utf-8")
    except Exception:
        return ""


def load_table_frame(att: dict) -> pd.DataFrame | None:
    """读取表格附件为 DataFrame（无索引列，保留用户原始列）"""
    if att["kind"] != "table":
        return None
    path = attachment_dir(att["session_id"]) / f"{att['id']}.table.csv"
    try:
        return pd.read_csv(path)
    except Exception as e:
        logger.warning(f"表格附件读取失败 {att['filename']}: {e}")
        return None


def build_manifest(att_rows: list[dict]) -> str:
    """构建注入 LLM 用户消息的附件清单（不落库，仅在发送给模型时拼接）"""
    if not att_rows:
        return ""
    lines = [
        (
            "[用户上传的附件]（全文用 read_attachment 工具按 offset 分页读取；"
            "表格类附件在 run_research_code 中已作为 attachments 变量 {文件名: DataFrame} 可直接使用）"
        )
    ]
    for a in att_rows:
        size_kb = a.get("size", 0) / 1024
        lines.append(f"- {a['filename']}（{'表格' if a['kind'] == 'table' else '文档'}，{size_kb:.0f}KB，id={a['id']}）")
        preview = (a.get("preview") or "").strip()
        if preview:
            lines.append(f"  预览: {preview[:800]}")
    return "\n".join(lines)


def delete_attachment_files(session_id: str, att_id: str) -> None:
    directory = attachment_dir(session_id)
    for path in directory.glob(f"{att_id}*"):
        try:
            path.unlink()
        except Exception:
            pass


def delete_session_attachments(session_id: str) -> None:
    """删除会话的全部附件文件（DB 行由会话删除事务一并清理）"""
    import shutil

    directory = attachment_dir(session_id)
    if directory.exists():
        shutil.rmtree(directory, ignore_errors=True)
